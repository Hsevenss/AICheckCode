#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Excel 处理模块
支持线程安全的读写操作
"""

import os
import threading
from typing import List, Dict, Any, Optional
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


class ExcelHandler:
    """Excel 处理器"""

    def __init__(self, file_path: str):
        """
        初始化 Excel 处理器

        Args:
            file_path: Excel 文件路径
        """
        self.file_path = os.path.abspath(file_path)
        self._lock = threading.Lock()
        self._workbook = None
        self._sheet = None
        self._load()

    def _load(self):
        """加载 Excel 文件"""
        self._workbook = load_workbook(self.file_path)
        self._sheet = self._workbook.active

    @property
    def sheet(self) -> Worksheet:
        """获取工作表"""
        return self._sheet

    @property
    def workbook(self):
        """获取工作簿"""
        return self._workbook

    def get_headers(self) -> List[str]:
        """获取表头"""
        headers = []
        for cell in self._sheet[1]:
            headers.append(str(cell.value) if cell.value else "")
        return headers

    def get_test_cases(self) -> List[Dict[str, Any]]:
        """
        获取所有测试用例（过滤空白行）

        Returns:
            测试用例列表，每个用例是一个字典
        """
        headers = self.get_headers()
        test_cases = []

        for row_idx, row in enumerate(self._sheet.iter_rows(min_row=2, values_only=True), start=2):
            # 过滤空白行：检查第一列（用例编号）是否有值
            if not row or len(row) == 0 or row[0] is None or str(row[0]).strip() == '':
                continue

            case = {}
            for col_idx, value in enumerate(row):
                if col_idx < len(headers):
                    case[headers[col_idx]] = value if value else ""
            case['_row'] = row_idx  # 记录行号用于写入结果
            test_cases.append(case)

        return test_cases

    def get_test_case(self, row: int) -> Dict[str, Any]:
        """
        获取指定行的测试用例

        Args:
            row: 行号（从 1 开始，1 为表头）

        Returns:
            测试用例字典
        """
        headers = self.get_headers()
        case = {}
        for col_idx, cell in enumerate(self._sheet[row]):
            if col_idx < len(headers):
                case[headers[col_idx]] = cell.value if cell.value else ""
        case['_row'] = row
        return case

    def write_result(self, row: int, test_result: str, validation_reason: str):
        """
        写入测试结果和校验依据（线程安全）

        Args:
            row: 行号（从 1 开始，1 为表头）
            test_result: 测试结果（通过/失败/异常）
            validation_reason: 校验理由
        """
        with self._lock:
            # H 列 = 第 8 列 = 测试结果
            # I 列 = 第 9 列 = 校验依据
            self._sheet[f'H{row}'] = test_result
            self._sheet[f'I{row}'] = validation_reason
            self._workbook.save(self.file_path)

    def write_single_result(self, row: int, column: str, value: str):
        """
        写入单个单元格结果（线程安全）

        Args:
            row: 行号
            column: 列名（如 'H', 'I'）
            value: 要写入的值
        """
        with self._lock:
            self._sheet[f'{column}{row}'] = value
            self._workbook.save(self.file_path)

    def get_column_index(self, column_name: str) -> int:
        """
        获取列名对应的索引

        Args:
            column_name: 列名

        Returns:
            列索引（从 0 开始），不存在返回 -1
        """
        headers = self.get_headers()
        try:
            return headers.index(column_name)
        except ValueError:
            return -1

    def close(self):
        """关闭 Excel 文件"""
        if self._workbook:
            self._workbook.close()
            self._workbook = None
            self._sheet = None


if __name__ == "__main__":
    # 测试 Excel 处理器
    import sys
    if len(sys.argv) > 1:
        handler = ExcelHandler(sys.argv[1])
        print("表头:", handler.get_headers())
        cases = handler.get_test_cases()
        print(f"测试用例数量：{len(cases)}")
        if cases:
            print("第一个用例:", cases[0])
        handler.close()
